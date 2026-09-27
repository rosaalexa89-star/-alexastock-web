import os
from datetime import timedelta, datetime
from functools import wraps
from io import BytesIO

from flask import Flask, render_template, request, jsonify, redirect, url_for, session, send_file
from openpyxl import Workbook
from openpyxl.styles import Font

import database

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "cambiar-esta-clave")
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=30)

USUARIO_VALIDO = os.environ.get("APP_USUARIO", "admin")
CLAVE_VALIDA = os.environ.get("APP_CLAVE", "admin")

# Se crean las tablas al arrancar la app (si ya existen, no hace nada).
with app.app_context():
    database.crear_tablas()


@app.before_request
def verificar_sesion():
    # Estas rutas quedan libres (sin sesión iniciada): la pantalla de
    # login en sí, y los archivos estáticos (CSS/JS si los hubiera).
    rutas_publicas = {"login"}

    if request.endpoint in rutas_publicas:
        return

    if request.endpoint and request.endpoint.startswith("static"):
        return

    if not session.get("logueado"):
        if request.path.startswith("/api/"):
            return jsonify({"ok": False, "mensaje": "Sesión expirada. Volvé a iniciar sesión."}), 401
        return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None

    if request.method == "POST":
        usuario = request.form.get("usuario", "")
        clave = request.form.get("clave", "")

        if usuario == USUARIO_VALIDO and clave == CLAVE_VALIDA:
            session.permanent = True
            session["logueado"] = True
            return redirect(url_for("index"))

        error = "Usuario o clave incorrectos."

    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ============================================================
# PÁGINA PRINCIPAL
# ============================================================

@app.route("/")
def index():
    return render_template("index.html")


# ============================================================
# PRODUCTOS
# ============================================================

@app.route("/api/productos", methods=["GET"])
def api_obtener_productos():
    return jsonify(database.obtener_productos())


@app.route("/api/productos", methods=["POST"])
def api_agregar_producto():
    datos = request.get_json(force=True)

    exito, mensaje, nuevo_id = database.agregar_producto(
        nombre=datos.get("nombre", ""),
        categoria=datos.get("categoria", ""),
        costo=float(datos.get("costo", 0) or 0),
        precio_venta=float(datos.get("precio_venta", 0) or 0),
        stock=int(datos.get("stock", 0) or 0),
        stock_minimo=int(datos.get("stock_minimo", 0) or 0),
    )

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje, "id": nuevo_id})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/productos/<int:producto_id>", methods=["PUT"])
def api_editar_producto(producto_id):
    datos = request.get_json(force=True)

    exito, mensaje = database.editar_producto(
        producto_id=producto_id,
        nombre=datos.get("nombre", ""),
        categoria=datos.get("categoria", ""),
        costo=float(datos.get("costo", 0) or 0),
        precio_venta=float(datos.get("precio_venta", 0) or 0),
        stock=int(datos.get("stock", 0) or 0),
        stock_minimo=int(datos.get("stock_minimo", 0) or 0),
    )

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/productos/<int:producto_id>", methods=["DELETE"])
def api_eliminar_producto(producto_id):
    exito, mensaje = database.eliminar_producto(producto_id)

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


# ============================================================
# CATEGORÍAS
# ============================================================

@app.route("/api/categorias", methods=["GET"])
def api_obtener_categorias():
    return jsonify(database.obtener_categorias())


@app.route("/api/categorias", methods=["POST"])
def api_agregar_categoria():
    datos = request.get_json(force=True)
    exito, mensaje = database.agregar_categoria(datos.get("nombre", ""))

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


# ============================================================
# MÉTODOS DE PAGO
# ============================================================

@app.route("/api/metodos_pago", methods=["GET"])
def api_obtener_metodos_pago():
    return jsonify(database.obtener_metodos_pago())


@app.route("/api/metodos_pago", methods=["POST"])
def api_agregar_metodo_pago():
    datos = request.get_json(force=True)
    exito, mensaje = database.agregar_metodo_pago(datos.get("nombre", ""))

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


# ============================================================
# VENTAS
# ============================================================

@app.route("/api/ventas", methods=["GET"])
def api_obtener_ventas():
    limite = request.args.get("limite", type=int)
    return jsonify(database.obtener_ventas(limite))


@app.route("/api/ventas", methods=["POST"])
def api_registrar_venta():
    datos = request.get_json(force=True)

    exito, mensaje = database.registrar_venta(
        producto_id=int(datos.get("producto_id")),
        cantidad=int(datos.get("cantidad", 0) or 0),
        precio_unitario=float(datos.get("precio_unitario", 0) or 0),
        medio_pago=datos.get("medio_pago", ""),
        cliente=datos.get("cliente", ""),
        pagado=float(datos.get("pagado", 0) or 0),
    )

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/ventas/<int:venta_id>/pago", methods=["POST"])
def api_registrar_pago(venta_id):
    datos = request.get_json(force=True)
    exito, mensaje = database.registrar_pago(venta_id, float(datos.get("monto", 0) or 0))

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/ventas/<int:venta_id>", methods=["PUT"])
def api_editar_venta(venta_id):
    datos = request.get_json(force=True)

    exito, mensaje = database.editar_venta(
        venta_id=venta_id,
        cantidad=int(datos.get("cantidad", 0) or 0),
        precio_unitario=float(datos.get("precio_unitario", 0) or 0),
        medio_pago=datos.get("medio_pago", ""),
        cliente=datos.get("cliente", ""),
        pagado=float(datos.get("pagado", 0) or 0),
    )

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/ventas/<int:venta_id>", methods=["DELETE"])
def api_eliminar_venta(venta_id):
    exito, mensaje = database.eliminar_venta(venta_id)

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


# ============================================================
# COMPRAS
# ============================================================

@app.route("/api/compras", methods=["GET"])
def api_obtener_compras():
    limite = request.args.get("limite", type=int)
    return jsonify(database.obtener_compras(limite))


@app.route("/api/compras", methods=["POST"])
def api_registrar_compra():
    datos = request.get_json(force=True)

    exito, mensaje = database.registrar_compra(
        producto_id=int(datos.get("producto_id")),
        cantidad=int(datos.get("cantidad", 0) or 0),
        moneda=datos.get("moneda", "ARS"),
        precio_unitario=float(datos.get("precio_unitario", 0) or 0),
        cotizacion=float(datos.get("cotizacion", 0) or 0),
        proveedor=datos.get("proveedor", ""),
    )

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/compras/<int:compra_id>", methods=["DELETE"])
def api_eliminar_compra(compra_id):
    exito, mensaje = database.eliminar_compra(compra_id)

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


# ============================================================
# REPORTE
# ============================================================

@app.route("/api/reporte", methods=["GET"])
def api_obtener_reporte():
    return jsonify(database.obtener_reporte())


# ============================================================
# GASTOS EXTRA
# ============================================================

@app.route("/api/gastos", methods=["GET"])
def api_obtener_gastos():
    return jsonify(database.obtener_gastos())


@app.route("/api/gastos", methods=["POST"])
def api_agregar_gasto():
    datos = request.get_json(force=True)

    exito, mensaje, nuevo_id = database.agregar_gasto(
        motivo=datos.get("motivo", ""),
        monto=float(datos.get("monto", 0) or 0),
    )

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje, "id": nuevo_id})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/gastos/<int:gasto_id>", methods=["DELETE"])
def api_eliminar_gasto(gasto_id):
    exito, mensaje = database.eliminar_gasto(gasto_id)

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


# ============================================================
# CONFIG (inversión inicial)
# ============================================================

@app.route("/api/config", methods=["GET"])
def api_obtener_config():
    return jsonify(database.obtener_config())


@app.route("/api/config", methods=["POST"])
def api_guardar_config():
    datos = request.get_json(force=True)

    exito, mensaje = database.guardar_inversion_inicial(
        float(datos.get("inversion_inicial", 0) or 0)
    )

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


# ============================================================
# PEDIDOS
# ============================================================

@app.route("/api/pedidos", methods=["GET"])
def api_obtener_pedidos():
    return jsonify(database.obtener_pedidos())


@app.route("/api/pedidos", methods=["POST"])
def api_agregar_pedido():
    datos = request.get_json(force=True)

    exito, mensaje, nuevo_id = database.agregar_pedido(
        producto_id=datos.get("producto_id"),
        cantidad=int(datos.get("cantidad", 0) or 0),
        cliente=datos.get("cliente", ""),
        notas=datos.get("notas", ""),
    )

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje, "id": nuevo_id})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/pedidos/<int:pedido_id>", methods=["DELETE"])
def api_eliminar_pedido(pedido_id):
    exito, mensaje = database.eliminar_pedido(pedido_id)

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/pedidos/<int:pedido_id>/entregar", methods=["POST"])
def api_entregar_pedido(pedido_id):
    datos = request.get_json(force=True)

    exito, mensaje = database.entregar_pedido(
        pedido_id=pedido_id,
        precio_unitario=float(datos.get("precio_unitario", 0) or 0),
        medio_pago=datos.get("medio_pago", ""),
        pagado=float(datos.get("pagado", 0) or 0),
    )

    if exito:
        return jsonify({"ok": True, "mensaje": mensaje})

    return jsonify({"ok": False, "mensaje": mensaje}), 400


@app.route("/api/exportar", methods=["GET"])
def api_exportar_excel():
    wb = Workbook()

    def fecha_txt(fecha):
        return fecha.strftime("%d/%m/%Y %H:%M") if fecha else ""

    # --- Ventas ---
    hoja_ventas = wb.active
    hoja_ventas.title = "Ventas"
    hoja_ventas.append([
        "Fecha", "Producto", "Cantidad", "Precio unitario", "Total",
        "Ganancia", "Medio de pago", "Cliente", "Pagado", "Saldo"
    ])
    for v in database.obtener_ventas():
        hoja_ventas.append([
            fecha_txt(v["fecha"]), v["producto_nombre"], v["cantidad"],
            v["precio_unitario"], v["total"], v["ganancia"], v["medio_pago"],
            v["cliente"], v["pagado"], v["saldo"]
        ])

    # --- Compras ---
    hoja_compras = wb.create_sheet("Compras")
    hoja_compras.append([
        "Fecha", "Producto", "Cantidad", "Moneda", "Costo unitario ($)",
        "Total ($)", "Proveedor"
    ])
    for c in database.obtener_compras():
        hoja_compras.append([
            fecha_txt(c["fecha"]), c["producto_nombre"], c["cantidad"], c["moneda"],
            c["costo_unitario_ars"], c["total_ars"], c["proveedor"]
        ])

    # --- Gastos extra ---
    hoja_gastos = wb.create_sheet("Gastos extra")
    hoja_gastos.append(["Fecha", "Motivo", "Monto"])
    for g in database.obtener_gastos():
        hoja_gastos.append([fecha_txt(g["fecha"]), g["motivo"], g["monto"]])

    # --- Pedidos pendientes ---
    hoja_pedidos = wb.create_sheet("Pedidos pendientes")
    hoja_pedidos.append(["Fecha", "Producto", "Cantidad", "Precio unitario", "Cliente", "Notas"])
    for p in database.obtener_pedidos():
        hoja_pedidos.append([
            fecha_txt(p["fecha"]), p["producto_nombre"], p["cantidad"],
            p["precio_unitario"], p["cliente"], p["notas"]
        ])

    # --- Inventario actual ---
    hoja_productos = wb.create_sheet("Inventario")
    hoja_productos.append(["Nombre", "Categoría", "Costo", "Precio de venta", "Stock", "Stock mínimo"])
    for p in database.obtener_productos():
        hoja_productos.append([
            p["nombre"], p["categoria"], p["costo"], p["precio_venta"], p["stock"], p["stock_minimo"]
        ])

    # Encabezados en negrita y primera fila congelada en cada hoja
    for hoja in wb.worksheets:
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        hoja.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"entre_casa_respaldo_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre_archivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


if __name__ == "__main__":
    app.run(debug=True)
